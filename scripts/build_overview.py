# Copyright 2026 思维刻度 SavorCode
# SPDX-License-Identifier: Apache-2.0
"""Build the Excel overview of the atlas (paradigm classes, concrete paradigms, markers, source checks).

Usage:  python scripts/build_overview.py [--out curation/reports/BCI范式图谱_vX.Y.Z_总览.xlsx] [--version X.Y.Z]
Requires openpyxl (pip install openpyxl). The workbook is a derived view; the YAML/CSV files remain the source of truth.
生成图谱的 Excel 总览（范式类、具体范式、标记物、出处核实）。总览只是派生视图，数据以 YAML/CSV 为准。
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from atlas import ROOT, load, read_yaml

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
PKG = {"perception": "A", "steady_state": "A", "motor": "B", "imagery": "B", "stimulation": "B",
       "error": "C", "control": "C", "memory": "D", "language": "D", "emotion": "E", "social": "E", "state": "E"}


def latest_version() -> str:
    m = re.search(r"^## \[(\d+\.\d+\.\d+)\]", (ROOT / "CHANGELOG.md").read_text(encoding="utf-8"), re.M)
    return m.group(1) if m else "dev"


def add_sheet(wb, title, rows, widths):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def joined(xs):
    return "; ".join(str(x) for x in xs or [])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--version", default=latest_version())
    ap.add_argument("--out")
    args = ap.parse_args()
    out = Path(args.out) if args.out else ROOT / "curation" / "reports" / f"BCI范式图谱_v{args.version}_总览.xlsx"

    atlas = load()
    classes = read_yaml(ROOT / "paradigms" / "_classes.yaml")["terms"]
    n_by_class: dict[str, int] = {}
    concrete = [["ID", "范式类", "包", "族", "中文名", "英文名", "状态", "BCI 类别", "记录模态", "刺激模态", "标记物",
                 "类别数", "类别", "提示/刺激", "刺激编码", "时序", "反馈", "区别说明", "源头文献", "DOI", "源头状态", "数据集"]]
    for pid, (_, d) in sorted(atlas.paradigms.items()):
        p, fs = d.get("protocol", {}) or {}, d.get("first_source", {}) or {}
        if d.get("status") != "deprecated":
            n_by_class[d["class"]] = n_by_class.get(d["class"], 0) + 1
        cit = fs.get("citation", "")
        concrete.append([pid, d["class"], PKG.get(d["family"], ""), d["family"], d["name"]["zh"], d["name"]["en"],
                         d.get("status", ""), joined(d.get("bci_category")), joined(d.get("recording_modality")),
                         joined(d.get("stimulus_modality")), joined(d.get("markers")), p.get("n_classes", ""),
                         joined(p.get("classes")), p.get("cue", ""), p.get("stimulus_coding", ""),
                         p.get("paradigm_timing", ""), p.get("feedback", ""), p.get("distinguishing", ""), cit,
                         fs.get("doi", ""), "TBD" if str(cit).startswith("TBD") else ("已核实" if fs.get("verified") else "待内容核实"),
                         "; ".join(x.get("name", "") for x in d.get("datasets", []) or [])])

    cls_rows = [["范式类", "包", "族", "中文名", "英文名", "具体范式数", "标记物", "类源头", "DOI"]]
    for c in classes:
        fs = c.get("first_source", {}) or {}
        cls_rows.append([c["id"], PKG.get(c["family"], ""), c["family"], c["name"]["zh"], c["name"]["en"],
                         n_by_class.get(c["id"], 0), joined(c.get("markers")), fs.get("citation", ""), fs.get("doi", "")])

    mk_rows = [["ID", "中文名", "英文名", "类型", "记录模态", "状态"]]
    for mid, (_, d) in sorted(atlas.markers.items()):
        mk_rows.append([mid, d["name"]["zh"], d["name"]["en"], d["type"], joined(d.get("recording_modality")), d.get("status", "")])

    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    active = sum(1 for _, d in atlas.paradigms.values() if d.get("status") != "deprecated")
    for r in [[f"BCI 脑机范式图谱 v{args.version} 总览"], [],
              ["范式类", len(classes)], ["具体范式（有效）", active], ["神经标记物", len(atlas.markers)], [],
              ["说明", "本表由 scripts/build_overview.py 从仓库数据生成，数据以 YAML/CSV 为准；出处已完成书目核实（D-072），内容核实待完成。"]]:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 100

    add_sheet(wb, "范式类", cls_rows, [14, 5, 13, 26, 36, 10, 36, 70, 28])
    add_sheet(wb, "具体范式", concrete, [16, 12, 5, 12, 26, 36, 10, 14, 18, 18, 30, 8, 36, 24, 20, 12, 10, 50, 70, 28, 12, 28])
    add_sheet(wb, "标记物", mk_rows, [28, 26, 40, 16, 24, 8])
    check = ROOT / "curation" / "verification" / "sources_check.csv"
    if check.exists():
        with open(check, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        add_sheet(wb, "出处核实", rows, [12] * len(rows[0]))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}: {len(classes)} classes, {active} concrete paradigms")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
